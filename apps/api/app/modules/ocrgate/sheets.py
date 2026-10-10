"""OCRGate input adapter: Excel / CSV exports -> a PDF with a real text layer.

Accounting systems and online banking often export .xlsx or .csv rather than PDF. Rendering the sheet as a PDF keeps one
code path for everything downstream: the same parsers, the page viewer and bounding-box evidence work without change,
and the officer sees exactly the cells that were read.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path

from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from ...datagen.render import _fonts

SHEET_EXT = {".xlsx", ".xlsm", ".csv"}
MAX_ROWS = 6000
MAX_COLS = 24
MAX_CELL = 110  # characters kept per cell; the column width uses the same cap so neighbouring cells never overlap
BALANCE_WORDS = ("остаток", "qoldiq", "qoldig", "balance", "saldo", "қолдиқ")


def _amount_text(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ")


def _load_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=False)
    ws = max(wb.worksheets, key=lambda w: w.max_row * min(w.max_column, MAX_COLS))  # the sheet that holds the table
    cells = [[c for c in row[:MAX_COLS]] for row in ws.iter_rows(max_row=MAX_ROWS)]
    ncols = max((len(r) for r in cells), default=0)
    # a column of numbers is a column of amounts when any cell has decimals or a 0.00 number format
    amount_cols = set()
    for j in range(ncols):
        for r in cells:
            c = r[j] if j < len(r) else None
            if c is not None and isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                if (isinstance(c.value, float) and not float(c.value).is_integer()) or "0.00" in (c.number_format or "") or "0,00" in (c.number_format or ""):
                    amount_cols.add(j)
                    break
    out: list[list[str]] = []
    for r in cells:
        row = []
        label = next((str(c.value).lower() for c in r if c.value not in (None, "")), "")
        balance_row = label.rstrip().endswith(":") and any(w in label for w in BALANCE_WORDS)
        for j, c in enumerate(r):
            v = c.value
            if v is None:
                row.append("")
            elif isinstance(v, (datetime, date)):
                row.append(v.strftime("%d.%m.%Y"))
            elif isinstance(v, bool):
                row.append(str(v))
            elif isinstance(v, (int, float)) and balance_row:  # "Входящий остаток:" | 1234567  -> an amount, not an id
                row.append(_amount_text(float(v)))
            elif isinstance(v, (int, float)):
                row.append(_amount_text(float(v)) if j in amount_cols and float(v) != 0 else ("" if j in amount_cols else str(int(v) if float(v).is_integer() else v)))
            else:
                row.append(str(v).strip())
        out.append(row)
    return out


def _load_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1")
    # csv.Sniffer picks the decimal comma in "1.234,56"; choose the delimiter whose cell count is most consistent across lines
    sample = [ln for ln in text.splitlines()[:60] if ln.strip()]
    best, best_score = ";", -1.0
    for d in (";", "\t", "|", ","):
        counts = [ln.count(d) for ln in sample]
        modal = max(set(counts), key=counts.count) if counts else 0
        score = (counts.count(modal) / len(counts)) * (1 if modal >= 2 else 0) + (0.01 if d == ";" else 0)
        if score > best_score:
            best, best_score = d, score

    class _D(csv.excel):
        delimiter = best
    return [[c.strip() for c in row[:MAX_COLS]] for _, row in zip(range(MAX_ROWS), csv.reader(io.StringIO(text), _D))]


def sheet_to_pdf(path: Path, out: Path | None = None) -> Path:
    """Render the sheet's non-empty rows, cells separated by wide gaps, to a landscape PDF (one row = one text line)."""
    rows = _load_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _load_csv(path)
    rows = [[c for c in r] for r in rows if any(c for c in r)]
    if not rows:
        raise ValueError("the spreadsheet is empty")
    out = out or path.with_suffix(path.suffix + ".pdf")
    _fonts()
    ncols = max(len(r) for r in rows)
    widths = [0] * ncols
    for r in rows:
        for j, c in enumerate(r):
            widths[j] = max(widths[j], min(len(c), MAX_CELL))
    W, H = landscape(A4)
    total_chars = sum(widths) + 3 * (ncols - 1)
    font = max(4.5, min(8.0, (W - 60) / max(total_chars, 1) / 0.6))
    gap = font * 0.6 * 3
    c = canvas.Canvas(str(out), pagesize=(W, H))
    y = H - 36
    line_h = font * 1.5
    page = 0
    for r in rows:
        if y < 36:
            c.showPage()
            y = H - 36
            page += 1
        c.setFont("DV", font)
        x = 30
        for j, cell in enumerate(r):
            if cell:
                c.drawString(x, y, cell[:MAX_CELL])
            x += widths[j] * font * 0.6 + gap
        y -= line_h
    c.save()
    return out
